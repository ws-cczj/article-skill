"""Check report structure and image presence; never certify scientific content."""
from __future__ import annotations
import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

SECTION_NAMES = {
    'background': ('研究背景',),
    'methods': ('研究方法',),
    'conclusions': ('主要结论', '研究结论', '关键图文解析', '主要结论与关键图文解析', '结论'),
    'innovations': ('创新点',),
    'citation': ('引用格式', '参考文献'),
}


def parenthetical_figure_reference(text: str) -> bool:
    # Standalone citation parentheses; retain normal subpanel labels such as 图2(a).
    return bool(re.search(r'[（(]\s*(?:见\s*)?(?:图|Fig(?:ure)?\.?)\s*\d+[^。；;\n）)]*[）)]',text,re.I))

@dataclass
class Paragraph:
    text: str
    style: str = ''
    images: int = 0
    missing_images: int = 0

@dataclass
class Block:
    title: str
    paragraphs: list[Paragraph]


def heading_level(p: Paragraph) -> int | None:
    match = re.fullmatch(r'(?:Heading\s*|标题\s*)([1-9])', p.style, re.I)
    if match:
        return int(match[1])
    match = re.match(r'^(#{1,6})\s+', p.text)
    return len(match[1]) if match else None


def clean_heading(text: str) -> str:
    text = re.sub(r'^#{1,6}\s+', '', text).strip().strip('*').strip()
    text = re.sub(r'^(?:[一二三四五六七八九十]+[、.．]|\d+[、.．])\s*', '', text)
    return text.rstrip('：:').strip()


def section_key(p: Paragraph) -> str | None:
    # Match the entire heading, never a keyword occurring inside ordinary prose.
    title = clean_heading(p.text)
    return next((key for key, names in SECTION_NAMES.items() if title in names), None)


def figure_ids(text: str) -> set[str]:
    ids = set()
    for kind, number in re.findall(r'(图|表|Fig(?:ure)?\.?|Table)[\s№#]*(\d+(?:[-–.]\d+)*)', text, re.I):
        prefix = 'table' if kind.lower() in ('表', 'table') else 'figure'
        ids.add(f'{prefix}:{number.replace(chr(8211), "-")}')
    return ids


def caption_ids(p: Paragraph) -> set[str]:
    if p.style and p.style not in ('Caption','题注'):
        return set()
    if re.match(r'^\s*(?:图|Fig(?:ure)?\.?)\s*\d+(?:[-–.]\d+)*(?=\s|[:：])', p.text, re.I):
        # A caption may mention other figures later; count its leading ID only.
        match = re.match(r'^\s*((?:图|Fig(?:ure)?\.?)\s*\d+(?:[-–.]\d+)*)', p.text, re.I)
        return figure_ids(match[1])
    return set()


def read_paragraphs(path: Path) -> list[Paragraph]:
    if path.suffix.lower() == '.docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph as WordParagraph
        doc = Document(path)
        paragraphs = []
        # Body paragraphs include DrawingML inline/anchor and legacy VML objects.
        for node in doc.element.body.xpath('.//w:p'):
            p = WordParagraph(node, doc._body)
            objects = p._p.xpath('.//a:blip | .//*[local-name()="imagedata" and namespace-uri()="urn:schemas-microsoft-com:vml"]')
            missing = 0
            for obj in objects:
                rid = obj.get(qn('r:embed')) or obj.get(qn('r:id')) or obj.get(qn('r:link'))
                rel = doc.part.rels.get(rid)
                if rel is None or rel.is_external:
                    missing += 1
                elif not rel.target_part.blob:
                    missing += 1
            paragraphs.append(Paragraph(p.text.strip(), p.style.name, len(objects), missing))
        return paragraphs
    if path.suffix.lower() not in ('.md', '.txt'):
        raise ValueError('仅支持 DOCX、Markdown 或纯文本；PDF请检查生成它的DOCX')
    result = []
    for chunk in re.split(r'\n\s*\n', path.read_text(encoding='utf-8')):
        # Separate adjacent headings and list items, retain wrapped prose as one paragraph.
        lines = chunk.strip().splitlines()
        if not lines:
            continue
        units = lines if any(re.match(r'^\s*(?:#{1,6}\s|[-*+]\s|\d+[.)、]\s*)', x) for x in lines) else [' '.join(lines)]
        for line in units:
            targets = re.findall(r'!\[[^\]]*\]\(([^)]+)\)', line)
            missing = 0
            for target in targets:
                target = target.strip().strip('<>')
                if re.match(r'^[a-z]+://', target, re.I) or not (path.parent / target).is_file():
                    missing += 1
            result.append(Paragraph(line.strip(), '', len(targets), missing))
    return result


def detect_blocks(paragraphs: list[Paragraph]) -> list[Block]:
    starts = [i for i, p in enumerate(paragraphs) if heading_level(p) == 1 and section_key(p) is None]
    if not starts:
        return [Block('document', paragraphs)] if paragraphs else []
    return [Block(paragraphs[i].text, paragraphs[i:starts[n+1] if n+1 < len(starts) else len(paragraphs)]) for n, i in enumerate(starts)]


def find_section_range(block: Block, names: tuple[str, ...]) -> tuple[int, int] | None:
    indices = [i for i, p in enumerate(block.paragraphs) if section_key(p)]
    for n, i in enumerate(indices):
        if clean_heading(block.paragraphs[i].text) in names:
            return i+1, indices[n+1] if n+1 < len(indices) else len(block.paragraphs)
    return None


def section_paragraphs(block: Block, key: str) -> list[Paragraph]:
    span = find_section_range(block, SECTION_NAMES[key])
    return block.paragraphs[span[0]:span[1]] if span else []


def section_text(block: Block, names: tuple[str, ...]) -> str:
    span = find_section_range(block, names)
    return '\n'.join(p.text for p in block.paragraphs[span[0]:span[1]]) if span else ''


def validate_block(block: Block, minimum_figures: int, innovation_count: int = 3) -> tuple[list[str], list[str]]:
    errors, warnings = [], []
    sections = [section_key(p) for p in block.paragraphs if section_key(p)]
    if sections != list(SECTION_NAMES):
        errors.append('五个章节必须各出现一次且顺序正确：研究背景、研究方法、主要结论、创新点、引用格式')
    title = next((p for p in block.paragraphs if p.text.strip()), None)
    if title is None or heading_level(title) != 1 or section_key(title):
        errors.append('每篇论文题目须为 Heading 1 或 Markdown # 标题')
    for p in block.paragraphs:
        level = heading_level(p)
        if section_key(p) and level != 2:
            errors.append(f'章节须用二级标题：{p.text}')
        elif level == 2 and not section_key(p):
            errors.append(f'方法/结论子标题不能与章节同级：{p.text}')
        elif level is not None and level > 3:
            warnings.append(f'请检查过深的标题层级：{p.text}')
        if p.missing_images:
            errors.append('存在缺失、空内容或未核实的外链图片')
    for key in SECTION_NAMES:
        if not any(p.text and heading_level(p) is None for p in section_paragraphs(block, key)):
            errors.append(f'章节无正文：{key}')
    figure_sequence=[]
    for p in block.paragraphs:
        if caption_ids(p):
            match=re.match(r'^\s*(?:图|Fig(?:ure)?\.?)\s*(\d+)',p.text,re.I)
            figure_sequence.append(int(match[1]))
    if figure_sequence != list(range(1,len(figure_sequence)+1)):
        errors.append('图号须按总结出现顺序从图1连续编号，方法与结论共用序列')
    bg = [p for p in section_paragraphs(block, 'background') if p.text and heading_level(p) is None]
    if len(bg) < 2:
        errors.append('研究背景少于2个正文段落')
    elif len(bg) > 3:
        warnings.append('研究背景超过通常的2–3段，请核对是否必要')
    conclusion = section_paragraphs(block, 'conclusions')
    figure_led=[p for p in conclusion if heading_level(p) is None and not caption_ids(p)
                and re.match(r'^\s*(?:图\s*\d+|Fig(?:ure)?\.?\s*\d+)',p.text,re.I)]
    if len(figure_led)>=2:
        warnings.append(f'主要结论有{len(figure_led)}段以图号开头，请结合论述需要判断是否成为读图清单；合理引用无需改写，不强制每段先写结论，也不能仅替换起句词')
    captions = set().union(*(caption_ids(p) for p in conclusion)) if conclusion else set()
    images = sum(p.images for p in conclusion)
    if len(captions) < minimum_figures:
        errors.append(f'主要结论仅有{len(captions)}个不同图注，要求至少{minimum_figures}个')
    if images < len(captions) or images < minimum_figures:
        errors.append(f'主要结论图片对象{images}个，图注{len(captions)}个；须核查真实图片，不能只写图号')
    for i, p in enumerate(conclusion):
        if caption_ids(p) and not any(x.images for x in conclusion[max(0,i-1):i+2]):
            errors.append(f'图注没有紧邻的图片：{p.text[:60]}')
    result_headings = [p for p in conclusion if heading_level(p) == 3]
    if not result_headings:
        errors.append('主要结论缺少三级结果小标题')
    if len(result_headings) < 5:
        errors.append(f'主要结论至少5项，检测到{len(result_headings)}项；按研究逻辑组织，不按图数或拆句凑条')
    for p in result_headings:
        if re.search(r'方法具有.*性质|方法(?:的)?(?:有效性|适用性|局限性)|方法边界', p.text):
            errors.append(f'方法评价不能独立充当主要结论：{p.text}')
    innovations = [p for p in section_paragraphs(block, 'innovations') if p.text and heading_level(p) is None]
    if len(innovations) != innovation_count:
        errors.append(f'创新点须为{innovation_count}个独立段落/条目，检测到{len(innovations)}个')
    citation = section_text(block, SECTION_NAMES['citation'])
    if not re.search(r'\b(?:19|20)\d{2}\b', citation):
        errors.append('引用缺少可识别年份')
    all_text = '\n'.join(p.text for p in block.paragraphs)
    if any(parenthetical_figure_reference(p.text) for p in block.paragraphs if not caption_ids(p)):
        errors.append('正文不要使用括号式图号引用（图X）；按需要自然写如图X所示，或省略图号')
    if re.search(r'以下将图表证据嵌入|图中数值和判断均保留论文|<TODO>|\{\{[^}]+\}\}', all_text):
        errors.append('正文包含执行说明或未替换占位符')
    return list(dict.fromkeys(errors)), list(dict.fromkeys(warnings))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', type=Path)
    parser.add_argument('--expected-papers', type=int)
    parser.add_argument('--min-figures', type=int, default=0, help='Optional explicit user figure quota; default has no figure count floor')
    parser.add_argument('--min-figures-per-paper', help='例如3,5；来自源文清点或用户要求')
    parser.add_argument('--innovation-count', type=int, default=3)
    args = parser.parse_args()
    try:
        if args.min_figures < 0 or args.innovation_count < 1 or (args.expected_papers is not None and args.expected_papers < 1):
            raise ValueError('篇数、创新点数须为正数；图数下限不得为负数')
        blocks = detect_blocks(read_paragraphs(args.path))
        minima = [int(x) for x in args.min_figures_per_paper.split(',')] if args.min_figures_per_paper else [args.min_figures]*len(blocks)
        if len(minima) != len(blocks) or any(x < 0 for x in minima):
            raise ValueError('逐篇图数下限须与检测篇数相同，且不得为负数')
    except Exception as exc:
        print(f'ERROR: {exc}')
        return 2
    errors, warnings = [], []
    if not blocks:
        errors.append('未检测到文献单元')
    if args.expected_papers is not None and len(blocks) != args.expected_papers:
        errors.append(f'期望{args.expected_papers}篇，检测到{len(blocks)}篇')
    for i, (block, minimum) in enumerate(zip(blocks, minima), 1):
        e, w = validate_block(block, minimum, args.innovation_count)
        errors.extend(f'文献{i}：{x}' for x in e)
        warnings.extend(f'文献{i}：{x}' for x in w)
    for w in warnings:
        print(f'WARNING: {w}')
    for e in errors:
        print(f'ERROR: {e}')
    print('人工仍须核验：学术标题质量、方法归类与精简、结果证据、图文真实性、引用准确性和逐页布局。')
    if errors:
        print('FAIL: 结构/图片检查未通过')
        return 1
    print(f'STRUCTURE PASS: {len(blocks)}篇结构与图片存在性通过；不代表事实及写作质量通过')
    return 0

if __name__ == '__main__':
    sys.exit(main())
