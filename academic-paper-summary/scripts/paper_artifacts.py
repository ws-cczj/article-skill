"""Reusable per-paper workspace, PDF cropping, page rendering and DOCX builder.

The agent supplies evidence-based content and reviewed crop coordinates in JSON.
This tool does not interpret papers or invent captions.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import shutil
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from uuid import uuid4


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


def inside(root, relative):
    path = (root / relative).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError('Path must remain inside the paper workspace')
    return path


def workspace(path):
    root = Path(path).resolve()
    if not (root/'manifest.json').is_file():
        raise ValueError('Not a paper workspace: manifest.json missing')
    return root


def unique_file(folder, stem, extension):
    for number in range(1, 10000):
        path = folder/f'{stem}-{number:03}{extension}'
        if not path.exists():
            return path
    raise ValueError('Too many output versions')


def initialize(pdf, output_root, name=None):
    import fitz
    source = Path(pdf).resolve()
    with fitz.open(source) as document:
        if document.needs_pass:
            raise ValueError('PDF is password protected')
        if len(document) == 0:
            raise ValueError('PDF has no pages')
        pages = [dict(page=i+1, width=p.rect.width, height=p.rect.height,
                      rotation=p.rotation, text=p.get_text()) for i,p in enumerate(document)]
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name or source.stem).strip(' .')[:65] or 'paper'
    root = Path(output_root).resolve()/f'{datetime.now():%Y%m%d_%H%M%S}_{stem}_{uuid4().hex[:8]}'
    root.mkdir(parents=True, exist_ok=False)
    for sub in ['source', 'assets', 'draft', 'review', 'final']:
        (root/sub).mkdir()
    shutil.copy2(source, root/'source/paper.pdf')
    save_json(root/'manifest.json',dict(schema_version=1, source_pdf='source/paper.pdf',original_name=source.name, page_count=len(pages)))
    save_json(root/'source/pages.json',pages)
    (root/'source/text.txt').write_text('\n\n'.join(f'=== Page {p["page"]} ===\n{p["text"]}' for p in pages),encoding='utf-8')
    save_json(root/'draft/report.json',dict(title_zh='',identity_images=[],background=[],methods=[],conclusions=[],innovations=[],citation='',citation_metadata={}))
    return root


def crop(root, page_number, box, label, dpi=240):
    import fitz
    if dpi < 72 or dpi > 600:
        raise ValueError('DPI must be between 72 and 600')
    if not re.fullmatch(r'[A-Za-z0-9_-]+',label):
        raise ValueError('Crop label must use ASCII letters, digits, underscores or hyphens')
    manifest = read_json(root/'manifest.json')
    with fitz.open(inside(root, manifest['source_pdf'])) as doc:
        if not 1 <= page_number <= len(doc):
            raise ValueError('Page number out of range (1-based)')
        page=doc[page_number-1]
        rect=fitz.Rect(box)
        if rect.is_empty or not page.rect.contains(rect):
            raise ValueError('Crop must be nonempty and inside the rendered page rectangle')
        output=unique_file(root/'assets',label,'.png')
        page.get_pixmap(matrix=fitz.Matrix(dpi/72,dpi/72), clip=rect, colorspace=fitz.csRGB, alpha=False).save(output)
    save_json(output.with_suffix('.json'),dict(source=manifest['source_pdf'],page=page_number,rect=list(box),dpi=dpi,image=output.relative_to(root).as_posix()))
    inspect_crop(root, output.relative_to(root).as_posix())
    return output


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def crop_review_path(root, image):
    folder = root/'review'/'crops'
    folder.mkdir(parents=True, exist_ok=True)
    key = hashlib.sha256(image.encode('utf-8')).hexdigest()[:16]
    return folder/f'{Path(image).stem}-{key}.json'


def inspect_crop(root, image):
    """Create a page-context preview and heuristic warnings; never approve a crop."""
    import fitz
    from PIL import Image, ImageDraw
    path = inside(root, image)
    meta = read_json(path.with_suffix('.json'))
    warnings = []
    with fitz.open(inside(root, meta['source'])) as doc:
        page = doc[meta['page']-1]
        rect = fitz.Rect(meta['rect'])
        if rect.is_empty or not page.rect.contains(rect):
            raise ValueError('Invalid crop provenance rectangle')
        # PDF text coordinates are unrotated; crop coordinates use the displayed page.
        for word in page.get_text('words'):
            box = fitz.Rect(word[:4]) * page.rotation_matrix
            if box.intersects(rect) and not rect.contains(box):
                warnings.append(dict(kind='cut_text', text=word[4], rect=list(box)))
        for block in page.get_text('dict')['blocks']:
            if block['type'] != 0:
                continue
            hits = []
            for line in block['lines']:
                box = fitz.Rect(line['bbox']) * page.rotation_matrix
                if not box.intersects(rect):
                    continue
                text = ''.join(s['text'] for s in line['spans']).strip()
                hits.append(text)
                if re.match(r'^(?:Fig(?:ure)?\.?\s*\d+|图\s*\d+)', text, re.I):
                    warnings.append(dict(kind='possible_source_caption', text=text, rect=list(box)))
            if len(hits) >= 3 and sum(len(t) for t in hits) >= 150:
                warnings.append(dict(kind='possible_body_text', text=' '.join(hits)[:300]))
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5,1.5), colorspace=fitz.csRGB, alpha=False)
        preview = Image.frombytes('RGB', (pix.width,pix.height), pix.samples)
        draw = ImageDraw.Draw(preview)
        sx,sy = pix.width/page.rect.width, pix.height/page.rect.height
        draw.rectangle((rect.x0*sx,rect.y0*sy,rect.x1*sx,rect.y1*sy), outline='red', width=3)
    review = crop_review_path(root,image)
    context = review.with_suffix('.png')
    preview.save(context)
    result = dict(image=image, sha256=digest(path), provenance_sha256=digest(path.with_suffix('.json')),
                  status='pending', context=context.relative_to(root).as_posix(), warnings=warnings,
                  note='View both the context preview and the actual crop. No warnings does not mean complete.')
    save_json(review,result)
    return review


def review_crop(root, image, note):
    """Record the agent's completed visual inspection, not a request for user approval."""
    if len(note.strip()) < 20:
        raise ValueError('Record concrete visual findings and resolve every warning (at least 20 characters)')
    path=inside(root,image)
    review=crop_review_path(root,image)
    if not review.exists():
        raise ValueError('Run inspect-crop, then view its context preview and actual crop first')
    result=read_json(review)
    if result['sha256'] != digest(path) or result['provenance_sha256'] != digest(path.with_suffix('.json')):
        raise ValueError('Crop or coordinates changed; run inspect-crop and inspect the new images again')
    result.update(status='reviewed', note=note.strip(), reviewed_at=datetime.now().isoformat())
    save_json(review,result)
    return review


def render_pages(root, pdf_relative, pages, dpi=110):
    import fitz
    if not 72 <= dpi <= 600:
        raise ValueError('DPI must be between 72 and 600')
    out=root/'review'/f'pages_{uuid4().hex[:8]}'
    with fitz.open(inside(root,pdf_relative)) as doc:
        indices=pages or list(range(1,len(doc)+1))
        if any(i<1 or i>len(doc) for i in indices):
            raise ValueError('Page number out of range (1-based)')
        out.mkdir()
        for n in indices:
            doc[n-1].get_pixmap(matrix=fitz.Matrix(dpi/72,dpi/72),colorspace=fitz.csRGB,alpha=False).save(out/f'page-{n:03}.png')
    return out


def validate_citation(data):
    """Cross-check citation against source metadata; source identity still needs visual verification."""
    meta = data.get('citation_metadata')
    if not isinstance(meta, dict):
        raise ValueError('citation_metadata is required: extract it from the current target paper')
    for field in ['original_title', 'journal', 'year']:
        if not isinstance(meta.get(field), str) or not meta[field].strip():
            raise ValueError(f'Current paper citation_metadata missing: {field}')
    if not isinstance(meta.get('authors'), list) or not meta['authors'] or not all(isinstance(a,str) and a.strip() for a in meta['authors']):
        raise ValueError('citation_metadata.authors must contain the current paper authors in source order')
    def normalized(text):
        return ''.join(c for c in unicodedata.normalize('NFKC',text).casefold() if c.isalnum())
    citation = data['citation']
    if normalized(meta['original_title']) not in normalized(citation):
        raise ValueError('Citation title does not match current paper original_title; do not reuse example citations')
    if not re.fullmatch(r'(?:19|20)\d{2}',meta['year']) or not re.search(r'(?<!\d)'+re.escape(meta['year'])+r'(?!\d)',citation):
        raise ValueError('Citation year does not match current paper metadata')
    doi = meta.get('doi')
    if doi:
        if not isinstance(doi,str):
            raise ValueError('citation_metadata.doi must be a string')
        doi = re.sub(r'^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)','',doi.strip(),flags=re.I)
        if doi.casefold() not in citation.casefold():
            raise ValueError('Citation must include the verified current paper DOI')


def validate_content(root, data):
    def nonempty(value): return isinstance(value,str) and bool(value.strip())
    if not nonempty(data.get('title_zh')) or not re.search(r'[\u3400-\u9fff]',data['title_zh']):
        raise ValueError('title_zh must contain an editable Chinese summary title')
    if not data.get('identity_images'):
        raise ValueError('English source title-area screenshot is required')
    if not 2 <= len(data.get('background',[])) <= 3:
        raise ValueError('background must contain 2 or 3 paragraphs')
    if len(data.get('innovations',[])) != 3 or not nonempty(data.get('citation')):
        raise ValueError('Three innovations and a citation are required')
    validate_citation(data)
    texts=[data['title_zh'],*data['background'],*data['innovations']]
    images=list(data['identity_images'])
    for section in ['methods','conclusions']:
        if not isinstance(data.get(section),list) or not data[section]:
            raise ValueError(f'{section} must contain at least one subsection')
        for item in data[section]:
            if not nonempty(item.get('heading')) or not item.get('paragraphs'):
                raise ValueError('Each subsection needs an unnumbered heading and paragraphs')
            if re.match(r'^\s*(?:\d+[.、]|[一二三四五]+、)',item['heading']):
                raise ValueError('Supply headings without numbering; the builder numbers each section from 1')
            texts += [item['heading'],*item['paragraphs']]
            for figure in item.get('figures',[]):
                if not nonempty(figure.get('caption')) or not isinstance(figure.get('number'),int) or figure['number']<1:
                    raise ValueError('Each figure needs a positive number, image and Chinese caption')
                texts.append(figure['caption']);images.append(figure['image'])
    if not all(nonempty(t) for t in texts):
        raise ValueError('All content paragraphs must be nonempty strings')
    for t in texts:
        if re.search(r'原文\s*(?:Fig(?:ure)?\.?|图)|作者(?:在|采用|利用|认为|构建)|本文(?:提出|研究)|该研究表明',t,re.I):
            raise ValueError('Rewrite indirect narration or source-caption prefixes before building')
    for name in images:
        path=inside(root,name)
        if not path.is_file() or path.suffix.lower() not in ('.png','.jpg','.jpeg'):
            raise ValueError(f'Image missing or unsupported: {name}')
        review=crop_review_path(root,name)
        if not review.exists():
            raise ValueError(f'Image needs inspect-crop and visual review: {name}')
        record=read_json(review)
        meta=path.with_suffix('.json')
        if (record.get('status') != 'reviewed' or record.get('sha256') != digest(path)
                or not meta.is_file() or record.get('provenance_sha256') != digest(meta)):
            raise ValueError(f'Image review missing or stale; inspect and review again: {name}')


def build(root, content='draft/report.json'):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image
    data=read_json(inside(root,content));validate_content(root,data)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2)
    sec.left_margin=sec.right_margin=Cm(2.2)
    specs={'Normal':('楷体',12,False),'Heading 1':('宋体',14,True),
           'Heading 2':('宋体',18,True),'Heading 3':('宋体',12,True),
           'Caption':('楷体',11,True),'Citation':('Times New Roman',12,False)}
    for name,(east,size,bold) in specs.items():
        style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,1)
        style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold
        style.font.color.rgb=RGBColor(0,0,0)
        rpr=style.element.get_or_add_rPr();fonts=rpr.rFonts
        fonts.set(qn('w:eastAsia'),east)
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+attr),None)
        fmt=style.paragraph_format
        fmt.keep_with_next=False;fmt.keep_together=False;fmt.page_break_before=False
        fmt.line_spacing=1.3;fmt.space_after=Pt(6)
    def paragraph(text,style='Normal'):
        p=doc.add_paragraph(text,style)
        p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False;p.paragraph_format.page_break_before=False
        if style=='Normal':p.paragraph_format.first_line_indent=Cm(.74)
        return p
    def picture(relative,max_height):
        path=inside(root,relative)
        with Image.open(path) as im: w,h=im.size
        width=min(16.6,max_height*w/h)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path),width=Cm(width))
    def figures(items):
        for f in items:
            picture(f['image'],17)
            paragraph(f'图{f["number"]} {f["caption"]}','Caption')
    paragraph(data['title_zh'],'Heading 1')
    for identity in data['identity_images']:picture(identity,8)
    paragraph('一、研究背景','Heading 2')
    for text in data['background']:paragraph(text)
    for key,title in [('methods','二、研究方法'),('conclusions','三、主要结论')]:
        paragraph(title,'Heading 2')
        for number,item in enumerate(data[key],1):
            p=paragraph(f'{number}. {item["heading"]}','Heading 3')
            if item.get('page_break_before'):
                # An explicit break is used instead of paragraph pagination flags.
                p.insert_paragraph_before().add_run().add_break(7)
            for text in item['paragraphs']:paragraph(text)
            figures(item.get('figures',[]))
    paragraph('四、创新点','Heading 2')
    for n,text in enumerate(data['innovations'],1):paragraph(f'{n}. {text}')
    paragraph('五、引用格式','Heading 2');paragraph(data['citation'],'Citation')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title=data['title_zh']
    output=unique_file(root/'final','summary','.docx');doc.save(output)
    return output


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    subs=parser.add_subparsers(dest='command',required=True)
    cmd=subs.add_parser('init');cmd.add_argument('--pdf',required=True);cmd.add_argument('--output-root',required=True);cmd.add_argument('--name')
    cmd=subs.add_parser('crop');cmd.add_argument('--workspace',required=True);cmd.add_argument('--page',type=int,required=True);cmd.add_argument('--rect',nargs=4,type=float,required=True);cmd.add_argument('--label',required=True);cmd.add_argument('--dpi',type=int,default=240)
    cmd=subs.add_parser('pages');cmd.add_argument('--workspace',required=True);cmd.add_argument('--pdf',default='source/paper.pdf');cmd.add_argument('--pages',nargs='+',type=int);cmd.add_argument('--dpi',type=int,default=110)
    cmd=subs.add_parser('inspect-crop');cmd.add_argument('--workspace',required=True);cmd.add_argument('--image',required=True)
    cmd=subs.add_parser('review-crop');cmd.add_argument('--workspace',required=True);cmd.add_argument('--image',required=True);cmd.add_argument('--note',required=True)
    cmd=subs.add_parser('build');cmd.add_argument('--workspace',required=True);cmd.add_argument('--content',default='draft/report.json')
    args=parser.parse_args()
    try:
        if args.command=='init': result=initialize(args.pdf,args.output_root,args.name)
        elif args.command=='crop':result=crop(workspace(args.workspace),args.page,args.rect,args.label,args.dpi)
        elif args.command=='pages':result=render_pages(workspace(args.workspace),args.pdf,args.pages,args.dpi)
        elif args.command=='inspect-crop':result=inspect_crop(workspace(args.workspace),args.image)
        elif args.command=='review-crop':result=review_crop(workspace(args.workspace),args.image,args.note)
        else:result=build(workspace(args.workspace),args.content)
        print(result)
    except Exception as exc:
        print(f'ERROR: {exc}',file=sys.stderr);return 1
    return 0

if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    raise SystemExit(main())
