"""Test artifact invariants with synthetic PDFs, not real paper summaries."""
import tempfile
import unittest
from pathlib import Path
import fitz
from docx import Document
from docx.oxml.ns import qn
from paper_artifacts import (initialize, crop, build, inside, save_json, read_json,
                            render_pages, inspect_crop, review_crop, crop_review_path, figure_mentions,
                            edge_warnings, reject_crop)


class ArtifactTests(unittest.TestCase):
    def test_caption_panel_coverage_and_real_descriptions(self):
        f=self.data['conclusions'][0]['figures'][0]
        f.update(panel_labels=['a','b'],caption='形貌对比。（a）HBA；（b）HCB。')
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'Panel descriptions'):build(self.root)
        f['caption']='（a）未处理组孔壁形貌，可见局部裂纹。'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'Missing caption'):build(self.root)
        f['caption']='（b）处理组孔壁形貌，裂纹范围较小；（a）未处理组孔壁形貌，可见局部裂纹。'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'order'):build(self.root)
        f['caption']='（a）未处理组孔壁形貌，可见局部裂纹；（b）处理组孔壁形貌，裂纹范围较小。'
        save_json(self.root/'draft/report.json',self.data)
        self.assertTrue(build(self.root).exists())

    def test_ordinary_check_preserves_delivery_receipt_and_failed_delivery_invalidates_it(self):
        from paper_memory import check
        output=build(self.root);name=output.relative_to(self.root).as_posix()
        check(self.root,name)
        stamp=self.root/'memory/last-delivery.json';before=stamp.read_bytes()
        check(self.root)
        self.assertEqual(stamp.read_bytes(),before)
        self.assertEqual(read_json(stamp)['artifact'],name)
        output.write_bytes(output.read_bytes()+b'changed')
        with self.assertRaises(ValueError):check(self.root,name)
        self.assertEqual(read_json(stamp)['status'],'artifact_check_pending')

    def test_parenthetical_figure_reference_rejected(self):
        from validate_summary import parenthetical_figure_reference
        for text in ['结果增加（图2）。','结果增加(图 2)。','结果增加（图2(a)、图3）。']:
            self.assertTrue(parenthetical_figure_reference(text))
        self.assertFalse(parenthetical_figure_reference('如图2(a)所示，结果增加。'))
        self.data['conclusions'][0]['paragraphs']=['结果增加（图2）。']
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'Parenthetical'):build(self.root)

    def test_identity_keywords_rejected(self):
        source=self.parent/'keywords.pdf'
        with fitz.open() as doc:
            p=doc.new_page();p.insert_text((40,50),'Title and authors')
            p.insert_text((40,80),'Affiliations and address')
            p.insert_text((40,110),'Keywords: composite; fatigue')
            doc.save(source)
        root=initialize(source,self.parent/'outputs')
        im=crop(root,1,[20,20,350,130],'header')
        name=im.relative_to(root).as_posix()
        review_crop(root,name,'Synthetic crop includes title, address and unwanted keyword line.')
        data=dict(self.data,identity_images=[name],methods=[dict(heading='方法',paragraphs=['方法内容'])],
                  conclusions=[dict(heading='结果',paragraphs=['结果内容'])],
                  methods_figure_absence_reason='Synthetic source has no figures.')
        save_json(root/'draft/report.json',data)
        with self.assertRaisesRegex(ValueError,'Keywords'):build(root)

    def test_final_check_rejects_wrong_or_changed_artifact(self):
        from paper_memory import check
        output=build(self.root)
        relative=output.relative_to(self.root).as_posix()
        self.data['background'][0]='内容已经修改'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'build inputs changed'):check(self.root,relative)
        output=build(self.root)
        output.write_bytes(output.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'build inputs changed'):
            check(self.root,output.relative_to(self.root).as_posix())
        with self.assertRaisesRegex(ValueError,'Missing build receipt'):
            check(self.root,'source/text.txt')

    def test_source_change_invalidates_crop_and_review(self):
        image=self.figure.relative_to(self.root).as_posix()
        source=self.root/'source/paper.pdf'
        source.write_bytes(source.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'stale'):build(self.root)
        with self.assertRaisesRegex(ValueError,'changed'):
            review_crop(self.root,image,'An old crop cannot be approved against a changed source.')
        with self.assertRaisesRegex(ValueError,'source changed'):inspect_crop(self.root,image)

    def test_schema_rejects_string_paragraphs_and_boolean_figure_number(self):
        self.data['methods'][0]['paragraphs']='必须是段落数组'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'paragraphs array'):build(self.root)
        self.data['methods'][0]['paragraphs']=['方法内容']
        self.data['methods'][0]['figures'][0]['number']=True
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'positive number'):build(self.root)

    def test_memory_final_check_binds_artifact(self):
        from paper_memory import check,sha
        output=build(self.root)
        result=check(self.root,output.relative_to(self.root).as_posix())
        self.assertEqual(result['status'],'memory_clear')
        self.assertEqual(result['artifact_sha256'],sha(output))
        self.assertEqual(result['content_sha256'],sha(self.root/'draft/report.json'))

    def test_crop_rejection_creates_memory_issue(self):
        from paper_memory import check
        reject_crop(self.root,self.figure.relative_to(self.root).as_posix(),'Synthetic cropped label requires a corrected source crop.')
        self.assertTrue(check(self.root)['blockers'])

    def test_raster_edge_ink_warning(self):
        from PIL import Image,ImageDraw
        path=self.parent/'edge.png'
        im=Image.new('RGB',(100,100),'white');ImageDraw.Draw(im).rectangle((0,30,20,50),fill='black');im.save(path)
        self.assertEqual([w['edge'] for w in edge_warnings(path)],['left'])

    def test_rejected_bytes_cannot_be_reapproved_or_renamed(self):
        image=self.figure.relative_to(self.root).as_posix()
        reject_crop(self.root,image,'Synthetic rejected image: known crop boundary defect.')
        with self.assertRaisesRegex(ValueError,'Rejected crop'):build(self.root)
        inspect_crop(self.root,image)
        with self.assertRaisesRegex(ValueError,'Rejected crop'):
            review_crop(self.root,image,'Attempting to approve identical rejected bytes should fail.')
        renamed=self.root/'assets/renamed.png';renamed.write_bytes(self.figure.read_bytes())
        renamed.with_suffix('.json').write_bytes(self.figure.with_suffix('.json').read_bytes())
        inspect_crop(self.root,'assets/renamed.png')
        with self.assertRaisesRegex(ValueError,'Rejected crop'):
            review_crop(self.root,'assets/renamed.png','Renaming rejected bytes cannot make them acceptable.')

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.parent=Path(self.tmp.name)
        source=self.parent/'input.pdf'
        with fitz.open() as doc:
            page=doc.new_page()
            page.insert_text((50,50),'English paper title and authors')
            page.draw_rect(fitz.Rect(50,100,250,200))
            page.insert_text((50,240),'Fig. 1. Synthetic caption outside crop')
            doc.save(source)
        self.root=initialize(source,self.parent/'outputs')
        self.identity=crop(self.root,1,[30,20,400,75],'identity')
        self.figure=crop(self.root,1,[40,90,260,210],'figure1')
        # Synthetic fixture attestations only: coordinates enclose the known text/rectangle.
        for image in [self.identity,self.figure]:
            review_crop(self.root,image.relative_to(self.root).as_posix(),
                        'Synthetic fixture: full known geometry enclosed; external caption and body excluded.')
        self.data=dict(title_zh='中文合成测试题目',identity_images=[self.identity.relative_to(self.root).as_posix()],
            methods_figure_absence_reason='',
            background=['背景一','背景二'],methods=[dict(heading='测试方法',paragraphs=['测试内容'],figures=[dict(number=1,image=self.figure.relative_to(self.root).as_posix(),caption='合成方法图')])],
            conclusions=[dict(heading='合成测试结果',paragraphs=['测试描述'],figures=[dict(number=2,image=self.figure.relative_to(self.root).as_posix(),caption='合成图注')])],
            innovations=['一项','二项','三项'],citation='Test A. English paper title and authors. Synthetic Journal, 2026.',
            citation_metadata=dict(original_title='English paper title and authors',authors=['Test A'],journal='Synthetic Journal',year='2026'))
        save_json(self.root/'draft/report.json',self.data)

    def test_workspace_and_version_isolation(self):
        second=initialize(self.root/'source/paper.pdf',self.parent/'outputs')
        self.assertNotEqual(self.root,second)
        a=build(self.root);original=a.read_bytes();b=build(self.root)
        self.assertNotEqual(a,b);self.assertEqual(a.read_bytes(),original)
        self.assertTrue(a.is_relative_to(self.root/'final'))

    def test_method_figure_or_source_absence_required(self):
        self.data['methods'][0]['figures']=[]
        self.data['conclusions'][0]['figures'][0]['number']=1
        self.data['methods_figure_absence_reason']='No dedicated method diagram; other source figures exist.'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'Methods need'):build(self.root)

    def test_truly_figureless_source_can_omit_methods_image(self):
        self.data['methods'][0]['figures']=[]
        self.data['conclusions'][0]['figures']=[]
        self.data['methods_figure_absence_reason']='Whole synthetic paper checked: no figures at all.'
        save_json(self.root/'draft/report.json',self.data)
        self.assertTrue(build(self.root).is_file())

    def test_figures_follow_report_order_across_sections(self):
        f=self.data['conclusions'][0]['figures'][0]
        self.data['methods'][0]['figures']=[dict(f,number=1,source_figure='Fig. 7')]
        f.update(number=2,source_figure='Fig. 3')
        del self.data['methods_figure_absence_reason']
        save_json(self.root/'draft/report.json',self.data)
        doc=Document(build(self.root))
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name=='Caption'],['图1 合成图注','图2 合成图注'])
        f['number']=3
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'appearance order'):build(self.root)

    def test_page_cache_reuses_and_invalidates(self):
        out=render_pages(self.root,'source/paper.pdf',[1])
        target=out/'page-001.png';timestamp=target.stat().st_mtime_ns
        self.assertEqual(render_pages(self.root,'source/paper.pdf',[1]),out)
        self.assertEqual(target.stat().st_mtime_ns,timestamp)
        expected=target.read_bytes();target.write_bytes(b'broken')
        render_pages(self.root,'source/paper.pdf',[1])
        self.assertEqual(target.read_bytes(),expected)
        self.assertNotEqual(render_pages(self.root,'source/paper.pdf',[1],dpi=120),out)
        with fitz.open() as doc:
            doc.new_page().insert_text((50,50),'Changed PDF')
            doc.save(self.root/'source/changed.pdf')
        self.assertNotEqual(render_pages(self.root,'source/changed.pdf',[1]),out)

    def test_figure_reference_index_is_candidate_evidence(self):
        items=figure_mentions([dict(page=3,text='See figure1a and Figure 1(b), Fig. 2c and Figs. 3a–c. Region ① is a location.')])
        self.assertEqual(len(items),4)
        self.assertTrue(all(x['page']==3 for x in items))
        self.assertEqual(items[0]['mention'],'figure1a')
        self.assertTrue((self.root/'source/figure_mentions.json').is_file())

    def test_chinese_title_and_english_screenshot_both_present(self):
        doc=Document(build(self.root))
        self.assertEqual(doc.paragraphs[0].text,'中文合成测试题目')
        self.assertTrue(doc.paragraphs[1]._p.xpath('.//a:blip'))
        self.assertEqual(doc.paragraphs[2].text,'一、研究背景')
        self.assertEqual(len(doc.inline_shapes),3)

    def test_numbering_fonts_and_no_paragraph_keep_flags(self):
        doc=Document(build(self.root))
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name=='Heading 3'],['1. 测试方法','1. 合成测试结果'])
        for name,east,size,bold in [('Normal','楷体',12,False),('Heading 1','宋体',14,True),('Heading 2','宋体',18,True),('Caption','楷体',11,True),('Citation','Times New Roman',12,False)]:
            style=doc.styles[name]
            self.assertEqual(style.font.size.pt,size)
            self.assertEqual(style.font.bold,bold)
            self.assertEqual(style.element.rPr.rFonts.get(qn('w:eastAsia')),east)
            self.assertEqual(style.element.rPr.rFonts.get(qn('w:ascii')),'Times New Roman')
            self.assertFalse(style.paragraph_format.keep_with_next)
            self.assertFalse(style.paragraph_format.keep_together)
            self.assertFalse(style.paragraph_format.page_break_before)

    def test_outside_paths_and_invalid_crop_rejected(self):
        with self.assertRaises(ValueError):inside(self.root,'../../escape.png')
        with self.assertRaises(ValueError):crop(self.root,1,[-1,0,30,40],'bad')
        self.data['identity_images']=['../other-paper/header.png']
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaises(ValueError):build(self.root)

    def test_missing_chinese_title_or_identity_rejected(self):
        self.data['title_zh']='English only';save_json(self.root/'draft/report.json',self.data)
        with self.assertRaises(ValueError):build(self.root)
        self.data['title_zh']='中文题目';self.data['identity_images']=[]
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaises(ValueError):build(self.root)

    def test_crop_provenance_and_page_preview(self):
        self.assertTrue(self.figure.with_suffix('.json').is_file())
        out=render_pages(self.root,'source/paper.pdf',[1])
        self.assertTrue((out/'page-001.png').is_file())

    def test_example_citation_cannot_replace_current_paper(self):
        self.data['citation']='Other A. A different example paper. Synthetic Journal, 2026.'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'title does not match'):build(self.root)

    def test_missing_citation_metadata_rejected(self):
        del self.data['citation_metadata']
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'citation_metadata'):build(self.root)

    def test_citation_year_and_doi_checked(self):
        self.data['citation_metadata']['year']='2025'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'year does not match'):build(self.root)
        self.data['citation_metadata']['year']='2026'
        self.data['citation_metadata']['doi']='10.0000/synthetic-test'
        save_json(self.root/'draft/report.json',self.data)
        with self.assertRaisesRegex(ValueError,'DOI'):build(self.root)
        self.data['citation']+=' https://doi.org/10.0000/synthetic-test'
        save_json(self.root/'draft/report.json',self.data)
        self.assertTrue(build(self.root).is_file())

    def test_crop_warns_on_caption_and_cut_text(self):
        image=crop(self.root,1,[60,35,260,250],'bad-crop').relative_to(self.root).as_posix()
        result=read_json(crop_review_path(self.root,image))
        self.assertEqual(result['status'],'pending')
        kinds={w['kind'] for w in result['warnings']}
        self.assertIn('cut_text',kinds)
        self.assertIn('possible_source_caption',kinds)
        self.assertTrue((self.root/result['context']).is_file())

    def test_build_blocks_pending_and_changed_images(self):
        image=self.figure.relative_to(self.root).as_posix()
        inspect_crop(self.root,image)
        with self.assertRaisesRegex(ValueError,'review missing'):build(self.root)
        review_crop(self.root,image,'Synthetic geometry checked: rectangle complete and caption excluded.')
        build(self.root)
        self.figure.write_bytes(self.figure.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'stale'):build(self.root)
        with self.assertRaisesRegex(ValueError,'changed'):review_crop(self.root,image,'This old inspection cannot apply to a changed image file.')

    def test_changed_coordinates_invalidate_review(self):
        meta=read_json(self.figure.with_suffix('.json'));meta['rect'][0]+=1
        save_json(self.figure.with_suffix('.json'),meta)
        with self.assertRaisesRegex(ValueError,'stale'):build(self.root)

    def test_rotated_page_text_warnings(self):
        source=self.parent/'rotated.pdf'
        with fitz.open() as doc:
            page=doc.new_page(width=300,height=400)
            page.insert_text((50,80),'Fig. 2. Caption')
            page.set_rotation(90)
            expected=fitz.Rect(page.get_text('words')[0][:4])*page.rotation_matrix
            doc.save(source)
        root=initialize(source,self.parent/'outputs')
        image=crop(root,1,[expected.x0-1,expected.y0+2,expected.x1+1,expected.y1+1],'rotated').relative_to(root).as_posix()
        kinds={w['kind'] for w in read_json(crop_review_path(root,image))['warnings']}
        self.assertIn('cut_text',kinds)
        self.assertIn('possible_source_caption',kinds)

    def test_neighboring_body_warning(self):
        source=self.parent/'columns.pdf'
        with fitz.open() as doc:
            page=doc.new_page()
            page.insert_textbox(fitz.Rect(310,50,550,300),
                'This is neighboring article body text. '*12,fontsize=11)
            doc.save(source)
        root=initialize(source,self.parent/'outputs')
        image=crop(root,1,[20,20,560,310],'wide').relative_to(root).as_posix()
        self.assertIn('possible_body_text',{w['kind'] for w in read_json(crop_review_path(root,image))['warnings']})


if __name__=='__main__':unittest.main(verbosity=2)
