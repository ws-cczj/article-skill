"""Synthetic structural regressions only; no target-paper summaries are generated."""
import base64
import copy
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from validate_summary import (Paragraph, Block, validate_block, figure_ids,
                              read_paragraphs, detect_blocks, section_text, SECTION_NAMES)


def fixture():
    return Block('测试结构', [
        Paragraph('测试结构', 'Heading 1'),
        Paragraph('一、研究背景', 'Heading 2'), Paragraph('背景一。'), Paragraph('背景二。'),
        Paragraph('二、研究方法', 'Heading 2'),
        Paragraph('1. 材料制备方法', 'Heading 3'), Paragraph('测试材料。'),
        Paragraph('三、主要结论', 'Heading 2'),
        Paragraph('正文含有结论、创新点和研究方法这些词。'),
        Paragraph('1. 测试对照存在差异', 'Heading 3'), Paragraph('测试证据。'),
        Paragraph('', images=1), Paragraph('图1 测试图注'),
        Paragraph('四、创新点', 'Heading 2'),
        Paragraph('测试条目一。'), Paragraph('测试条目二。'), Paragraph('测试条目三。'),
        Paragraph('五、引用格式', 'Heading 2'), Paragraph('测试引用，2020。')])

class ValidatorTests(unittest.TestCase):
    def test_source_numbering_cannot_replace_report_sequence(self):
        b=fixture();b.paragraphs[12].text='图7 测试图注'
        self.assertTrue(any('连续编号' in x for x in validate_block(b,1)[0]))

    def test_valid_structure(self):
        self.assertEqual(validate_block(fixture(), 1)[0], [])

    def test_keywords_in_prose_do_not_end_section(self):
        text = section_text(fixture(), SECTION_NAMES['conclusions'])
        self.assertIn('图1', text)
        self.assertIn('创新点', text)

    def test_figure_and_table_ids_are_distinct(self):
        self.assertEqual(figure_ids('图1 表1 Fig. 1 Table 1 图2'), {'figure:1','table:1','figure:2'})

    def test_caption_without_image_fails(self):
        b = fixture(); b.paragraphs[11].images = 0
        self.assertTrue(any('图片' in x for x in validate_block(b,1)[0]))

    def test_repeated_figure_does_not_meet_count(self):
        b = fixture(); b.paragraphs[13:13] = [Paragraph('',images=1), Paragraph('图1 重复图注')]
        self.assertTrue(any('不同图注' in x for x in validate_block(b,2)[0]))

    def test_wrong_subheading_level_fails(self):
        b = fixture(); b.paragraphs[5].style='Heading 2'
        self.assertTrue(any('同级' in x for x in validate_block(b,1)[0]))

    def test_extra_innovation_fails(self):
        b = fixture(); b.paragraphs.insert(17,Paragraph('第四条。'))
        self.assertTrue(any('创新点须' in x for x in validate_block(b,1)[0]))

    def test_empty_citation_and_out_of_order_fail(self):
        b=fixture();b.paragraphs[-1].text=''
        self.assertTrue(any('citation' in x for x in validate_block(b,1)[0]))
        b=fixture();b.paragraphs[1],b.paragraphs[4]=b.paragraphs[4],b.paragraphs[1]
        self.assertTrue(any('顺序' in x for x in validate_block(b,1)[0]))

    def test_method_assessment_is_flagged(self):
        b=fixture();b.paragraphs[9].text='1. 方法具有半定量性质'
        self.assertTrue(any('方法评价' in x for x in validate_block(b,1)[0]))

    def test_multiple_papers_keep_figures_separate(self):
        p=fixture().paragraphs+copy.deepcopy(fixture().paragraphs)
        blocks=detect_blocks(p);self.assertEqual(len(blocks),2)
        blocks[1].paragraphs[11].images=0
        self.assertEqual(validate_block(blocks[0],1)[0],[])
        self.assertTrue(validate_block(blocks[1],1)[0])

    def test_docx_images_inline_anchor_vml_and_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp); png=path/'pixel.png'
            png.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='))
            doc=Document();p=doc.add_paragraph();shape=p.add_run().add_picture(str(png))
            rid=shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
            p2=doc.add_paragraph();shape2=p2.add_run().add_picture(str(png))
            shape2._inline.tag=qn('wp:anchor') # Synthetic anchor; parser test, not layout delivery.
            from lxml import etree
            legacy=etree.Element('{urn:schemas-microsoft-com:vml}imagedata');legacy.set(qn('r:id'),rid)
            doc.add_paragraph()._p.append(legacy)
            doc.add_table(rows=1,cols=1).cell(0,0).paragraphs[0].add_run().add_picture(str(png))
            out=path/'test.docx';doc.save(out)
            result=read_paragraphs(out)
            self.assertEqual(sum(p.images for p in result),4)
            self.assertEqual(sum(p.missing_images for p in result),0)

    def test_markdown_missing_image_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'test.md';path.write_text('# Test\n\n![test](missing.png)',encoding='utf-8')
            self.assertEqual(sum(p.missing_images for p in read_paragraphs(path)),1)

if __name__=='__main__':
    unittest.main(verbosity=2)
