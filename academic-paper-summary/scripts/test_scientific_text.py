import io
import unittest
from docx import Document
from docx.oxml.ns import qn
from scientific_text import scientific_runs, add_scientific_text


class ScientificTextTests(unittest.TestCase):
    def test_explicit_exponents(self):
        self.assertEqual(list(scientific_runs('3×10^4次，10^{-3}，5×10⁴')),
                         [('3×10', None), ('4', 'superscript'), ('次，10', None),
                          ('-3', 'superscript'), ('，5×10', None), ('4', 'superscript')])

    def test_no_inference(self):
        text = 'CO2 T700 HE-3 2026 5×104 图6 10^4sample 10^4.5'
        self.assertEqual(list(scientific_runs(text)), [(text, None)])

    def test_native_word_roundtrip_and_style(self):
        doc = Document()
        doc.styles['Caption'].font.bold = True
        p = doc.add_paragraph(style='Caption')
        add_scientific_text(p, 'CO₂、SO₄²⁻、mm³、3×10^4次')
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        result = Document(buffer).paragraphs[0]
        self.assertEqual(result.text, 'CO2、SO42-、mm3、3×104次')
        aligned = [(r.text, r._r.rPr.find(qn('w:vertAlign')).get(qn('w:val')))
                   for r in result.runs if r._r.rPr is not None]
        self.assertEqual(aligned, [('2', 'subscript'), ('4', 'subscript'),
                                  ('2-', 'superscript'), ('3', 'superscript'),
                                  ('4', 'superscript')])
        self.assertTrue(result.style.font.bold)
        self.assertTrue(all(r.bold is None for r in result.runs))


if __name__ == '__main__':
    unittest.main()
