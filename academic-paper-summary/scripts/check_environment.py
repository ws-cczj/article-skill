"""Read-only preflight. --smoke uses auto-cleaned synthetic files, never papers."""
from __future__ import annotations

import argparse
import importlib
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


PACKAGES = {"docx": "python-docx", "fitz": "PyMuPDF", "PIL": "Pillow"}


def package_checks() -> list[dict]:
    checks = []
    for module, package in PACKAGES.items():
        try:
            loaded = importlib.import_module(module)
            version = getattr(loaded, "__version__", getattr(loaded, "VersionBind", "unknown"))
            checks.append(dict(name=package, status="OK", detail=str(version)))
        except Exception as exc:
            checks.append(dict(name=package, status="MISSING", detail=f"{type(exc).__name__}: {exc}"))
    return checks


def renderer_checks() -> list[dict]:
    candidates = [shutil.which("soffice"), shutil.which("libreoffice")]
    if sys.platform == "win32":
        for key in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
            if os.environ.get(key):
                candidates.append(str(Path(os.environ[key]) / "LibreOffice/program/soffice.exe"))
    elif sys.platform == "darwin":
        candidates.append("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return [dict(name="DOCX renderer", status="DETECTED", detail=candidate)]
    if sys.platform == "win32":
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CLSID"):
                return [dict(name="DOCX renderer", status="DETECTED", detail="Microsoft Word COM registration; actual export still needs verification")]
        except OSError:
            pass
    return [dict(name="DOCX renderer", status="UNKNOWN", detail="No automatic Word/LibreOffice route detected; configure a DOCX renderer or use the agent's available document tools")]


def font_checks() -> list[dict]:
    names: set[str] = set()
    if sys.platform == "win32":
        import winreg
        key_name = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                with winreg.OpenKey(hive, key_name) as key:
                    for i in range(winreg.QueryInfoKey(key)[1]):
                        name, value, _ = winreg.EnumValue(key, i)
                        names.add(name.casefold())
                        names.add(Path(str(value)).name.casefold())
            except OSError:
                continue
        simsun = any(n == "simsun.ttc" or n.startswith("simsun ") or n.startswith("simsun &") or n.startswith("宋体 ") for n in names)
        times = any(n == "times.ttf" or n.startswith("times new roman (") for n in names)
    else:
        command = shutil.which("fc-list")
        if command:
            try:
                result = subprocess.run([command, "--format=%{family}\n"], capture_output=True, text=True, timeout=15, check=True)
                names = {name.strip().casefold() for line in result.stdout.splitlines() for name in line.split(",")}
            except (OSError, subprocess.SubprocessError):
                pass
        simsun = bool(names & {"simsun", "宋体"})
        times = "times new roman" in names
    return [dict(name=name, status="DETECTED" if found else "UNKNOWN", detail="Font name detected; confirm actual document rendering" if found else "Not confirmed. Install a properly licensed font or verify via the renderer; do not silently substitute") for name, found in (("SimSun / 宋体", simsun), ("Times New Roman", times))]


def smoke_check() -> dict:
    """Exercise PDF extraction/rasterization and DOCX image embedding, not Word rendering."""
    try:
        import fitz
        from docx import Document
        from PIL import Image

        with tempfile.TemporaryDirectory(prefix="article-skill-check-") as tmp:
            root = Path(tmp)
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "article-skill environment check")
            pdf.save(root / "test.pdf")
            pdf.close()
            with fitz.open(root / "test.pdf") as source:
                if "environment check" not in source[0].get_text():
                    raise RuntimeError("PDF text extraction failed")
                source[0].get_pixmap().save(root / "test.png")
            with Image.open(root / "test.png") as picture:
                picture.verify()
            doc = Document()
            doc.add_paragraph("Synthetic environment check")
            doc.add_picture(str(root / "test.png"))
            doc.save(root / "test.docx")
            if len(Document(root / "test.docx").inline_shapes) != 1:
                raise RuntimeError("DOCX image embedding failed")
        return dict(name="Document libraries smoke test", status="OK", detail="PDF text/raster and DOCX image round-trip passed; DOCX rendering and font fidelity NOT tested")
    except Exception as exc:
        return dict(name="Document libraries smoke test", status="MISSING", detail=f"{type(exc).__name__}: {exc}")


def exit_status(checks: list[dict]) -> int:
    if any(c["status"] == "MISSING" for c in checks):
        return 1
    if any(c["status"] == "UNKNOWN" for c in checks):
        return 2
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--json", action="store_true", help="machine-readable report")
    parser.add_argument("--smoke", action="store_true", help="exercise libraries using temporary synthetic documents")
    args = parser.parse_args()
    checks = [dict(name="Python", status="OK" if sys.version_info >= (3, 10) else "MISSING", detail=f"{platform.python_version()} at {sys.executable}; recommended minimum 3.10")]
    checks += package_checks() + renderer_checks() + font_checks()
    if args.smoke:
        checks.append(smoke_check())
    code = exit_status(checks)
    report = dict(exit_code=code, checks=checks, note="Detection only; actual DOCX export, exact fonts and scientific quality require separate verification. No packages installed and no welcome state changed.")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        for check in checks:
            print(f'[{check["status"]}] {check["name"]}: {check["detail"]}')
        print(report["note"])
        if code:
            print("Setup guidance: references/environment-setup.md (inside the skill folder)")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
