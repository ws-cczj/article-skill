import io
import unittest
from datetime import datetime, timezone
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from document_properties import prepare_report_properties


class DocumentPropertiesTests(unittest.TestCase):
    def test_clean_package_preserves_content(self):
        doc = Document()
        doc.add_paragraph('正文 CO2 与 3×10⁴').add_run('保留格式').bold = True
        original_xml = doc.part.blob
        before = datetime.now(timezone.utc).replace(microsecond=0)
        prepare_report_properties(doc, '文献总结')
        output = io.BytesIO()
        doc.save(output)
        with ZipFile(output) as package:
            self.assertIsNone(package.testzip())
            self.assertEqual(package.read('word/document.xml'), original_xml)
            self.assertFalse(any('thumbnail' in n for n in package.namelist()))
            self.assertNotIn(b'metadata/thumbnail', package.read('_rels/.rels'))
            self.assertNotIn(b'thumbnail', package.read('[Content_Types].xml'))
            app = ET.fromstring(package.read('docProps/app.xml'))
            self.assertEqual([e.text for e in app], ['article-skill'])
        output.seek(0)
        result = Document(output)
        properties = result.core_properties
        self.assertEqual(properties.title, '文献总结')
        self.assertEqual(properties.author, '捶捶自己')
        self.assertEqual(properties.comments, '')
        created = properties.created.replace(tzinfo=timezone.utc)
        self.assertGreaterEqual(created, before)
        self.assertEqual(properties.created, properties.modified)
        self.assertTrue(result.paragraphs[0].runs[-1].bold)


if __name__ == '__main__':
    unittest.main()
